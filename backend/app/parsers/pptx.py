import io
from typing import List, Dict, Any, Union
from pptx import Presentation

class PptxParser:
    """Parses PowerPoint .pptx presentations into slide text chunks."""

    @staticmethod
    def parse(file_or_path: Union[str, bytes], filename: str = "presentation.pptx") -> List[Dict[str, Any]]:
        chunks = []
        try:
            if isinstance(file_or_path, bytes):
                prs = Presentation(io.BytesIO(file_or_path))
            else:
                prs = Presentation(file_or_path)

            for slide_idx, slide in enumerate(prs.slides):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())

                if slide_texts:
                    content = "\n".join(slide_texts)
                    chunks.append({
                        "source_file": f"{filename} (Slide {slide_idx + 1})",
                        "content": content,
                        "span": content[:120] + ("..." if len(content) > 120 else "")
                    })
        except Exception as e:
            chunks.append({
                "source_file": filename,
                "content": f"Error parsing PPTX: {str(e)}",
                "span": "Error reading PPTX file"
            })
        return chunks
