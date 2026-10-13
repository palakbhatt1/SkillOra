import io
from typing import List, Dict, Any, Union
from docx import Document

class DocxParser:
    """Parses Word .docx documents into text chunks."""

    @staticmethod
    def parse(file_or_path: Union[str, bytes], filename: str = "document.docx") -> List[Dict[str, Any]]:
        chunks = []
        try:
            if isinstance(file_or_path, bytes):
                doc = Document(io.BytesIO(file_or_path))
            else:
                doc = Document(file_or_path)

            for p_idx, paragraph in enumerate(doc.paragraphs):
                text = paragraph.text.strip()
                if text:
                    chunks.append({
                        "source_file": filename,
                        "content": text,
                        "span": text[:120] + ("..." if len(text) > 120 else "")
                    })

            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        chunks.append({
                            "source_file": filename,
                            "content": row_text,
                            "span": row_text[:120] + ("..." if len(row_text) > 120 else "")
                        })
        except Exception as e:
            chunks.append({
                "source_file": filename,
                "content": f"Error parsing DOCX: {str(e)}",
                "span": "Error reading DOCX file"
            })
        return chunks
